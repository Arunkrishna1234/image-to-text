from flask import Flask, render_template, request, jsonify, send_file, session
from werkzeug.utils import secure_filename
import os
import cv2
import pytesseract
import numpy as np
import re
from datetime import datetime
import fitz  # PyMuPDF
from docx import Document
import logging
from functools import wraps
import uuid
import json
from PIL import Image
import io
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# Initialize Flask app
app = Flask(__name__)
app.config['SECRET_KEY'] = 'your-secret-key-here'  # Change this!
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['OUTPUT_FOLDER'] = 'clean_notes'
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max file size
app.config['ALLOWED_EXTENSIONS'] = {'png', 'jpg', 'jpeg', 'bmp', 'tiff', 'pdf', 'docx'}

# Setup logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Create necessary folders
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['OUTPUT_FOLDER'], exist_ok=True)

# =======================
# UTILITY FUNCTIONS
# =======================

def allowed_file(filename):
    """Check if file extension is allowed"""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def generate_unique_filename(original_filename):
    """Generate unique filename to prevent collisions"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    unique_id = str(uuid.uuid4())[:8]
    name, ext = os.path.splitext(secure_filename(original_filename))
    return f"{name}_{timestamp}_{unique_id}{ext}"

# =======================
# TEXT PROCESSING
# =======================

def smart_text_cleaner(text, preserve_newlines=False):
    """Enhanced text cleaning with options"""
    if not text:
        return ""
    
    lines = text.split('\n')
    clean_lines = []
    
    for line in lines:
        stripped = line.strip()
        
        if len(stripped) <= 1:
            continue
        
        alpha_chars = sum(1 for c in stripped if c.isalpha())
        digit_chars = sum(1 for c in stripped if c.isdigit())
        total_chars = len(stripped)
        
        alpha_digit_ratio = (alpha_chars + digit_chars) / total_chars if total_chars > 0 else 0
        
        if alpha_digit_ratio > 0.3:
            line_clean = re.sub(r'[^\w\s\.,!?;:()\-]', '', stripped)
            line_clean = re.sub(r'\s+', ' ', line_clean).strip()
            
            if line_clean and len(line_clean) > 2:
                clean_lines.append(line_clean)
    
    separator = '\n' if preserve_newlines else ' '
    return separator.join(clean_lines)

def enhanced_preprocess(img):
    """Advanced image preprocessing for better OCR"""
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    height, width = gray.shape
    
    # Upscale small images
    if height < 100 or width < 200:
        scale = 2
        gray = cv2.resize(gray, (width * scale, height * scale), 
                         interpolation=cv2.INTER_CUBIC)
    
    # Denoise
    try:
        denoised = cv2.fastNlMeansDenoising(gray, None, 10, 7, 21)
    except:
        denoised = gray
    
    # Multiple thresholding techniques
    blurred = cv2.GaussianBlur(denoised, (3, 3), 0)
    
    _, thresh_otsu = cv2.threshold(blurred, 0, 255, 
                                   cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    
    thresh_adapt = cv2.adaptiveThreshold(blurred, 255, 
                                         cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                         cv2.THRESH_BINARY, 11, 2)
    
    kernel = np.ones((2, 2), np.uint8)
    thresh_morph = cv2.morphologyEx(thresh_otsu, cv2.MORPH_CLOSE, kernel)
    
    return [thresh_otsu, thresh_adapt, thresh_morph, denoised]

# =======================
# EXTRACTION FUNCTIONS
# =======================

def extract_text_from_image(image_path):
    """Extract text from image with multiple OCR attempts"""
    try:
        image = cv2.imread(image_path)
        if image is None:
            return None, "Could not load image"
        
        thresh_images = enhanced_preprocess(image)
        all_texts = []
        psm_configs = [
            '--oem 3 --psm 6',
            '--oem 3 --psm 3',
            '--oem 3 --psm 4',
            '--oem 3 --psm 11'
        ]
        
        for processed_img in thresh_images:
            for config in psm_configs:
                try:
                    text = pytesseract.image_to_string(processed_img, config=config)
                    cleaned = smart_text_cleaner(text)
                    if cleaned and len(cleaned) > 10:
                        all_texts.append(cleaned)
                except Exception as e:
                    logger.warning(f"OCR attempt failed: {e}")
                    continue
        
        if not all_texts:
            return None, "No readable text found in image"
        
        # Return the longest extracted text
        best_text = max(all_texts, key=len)
        return best_text, None
        
    except Exception as e:
        logger.error(f"Image extraction error: {e}")
        return None, str(e)

def extract_text_from_pdf(pdf_path):
    """Extract text from PDF with improved error handling"""
    try:
        text = ""
        
        # Method 1: Try direct text extraction with PyMuPDF
        try:
            logger.info("Attempting direct text extraction from PDF...")
            with fitz.open(pdf_path) as doc:
                for page_num, page in enumerate(doc):
                    page_text = page.get_text("text")
                    if page_text.strip():
                        text += page_text + "\n"
                    logger.info(f"Extracted from page {page_num + 1}: {len(page_text)} chars")
        except Exception as e:
            logger.warning(f"Direct text extraction failed: {e}")
        
        # Method 2: If no text found, convert to images and OCR
        if not text.strip():
            logger.info("No direct text found. Converting PDF to images for OCR...")
            try:
                # Use PyMuPDF to render pages as images (more reliable than pdf2image)
                with fitz.open(pdf_path) as doc:
                    for page_num, page in enumerate(doc):
                        # Render page to image
                        pix = page.get_pixmap(matrix=fitz.Matrix(300/72, 300/72))  # 300 DPI
                        img_data = pix.tobytes("png")
                        
                        # Convert to PIL Image
                        pil_image = Image.open(io.BytesIO(img_data))
                        
                        # Convert PIL to OpenCV format
                        opencv_image = cv2.cvtColor(np.array(pil_image), cv2.COLOR_RGB2BGR)
                        
                        # Preprocess and OCR
                        processed_images = enhanced_preprocess(opencv_image)
                        
                        for proc_img in processed_images:
                            try:
                                page_text = pytesseract.image_to_string(proc_img, config='--oem 3 --psm 6')
                                if page_text.strip():
                                    text += page_text + "\n"
                                    break  # Use first successful extraction
                            except Exception as e:
                                logger.warning(f"OCR failed on page {page_num + 1}: {e}")
                                continue
                        
                        logger.info(f"OCR completed for page {page_num + 1}")
                        
            except Exception as e:
                logger.error(f"PDF to image conversion failed: {e}")
                return None, f"Failed to process PDF: {str(e)}"
        
        # Clean the extracted text
        if text.strip():
            cleaned_text = smart_text_cleaner(text, preserve_newlines=True)
            
            if not cleaned_text:
                return None, "No readable text found in PDF after cleaning"
            
            return cleaned_text, None
        else:
            return None, "No text could be extracted from PDF. The file may be empty or contain only images without recognizable text."
        
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return None, f"Error processing PDF: {str(e)}"

def extract_text_from_docx(docx_path):
    """Extract text from Word document"""
    try:
        doc = Document(docx_path)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += "\n" + cell.text
        
        cleaned_text = smart_text_cleaner(text, preserve_newlines=True)
        
        if not cleaned_text:
            return None, "No readable text found in document"
        
        return cleaned_text, None
        
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return None, str(e)

def extract_text_from_file(file_path):
    """Main extraction function - routes to appropriate handler"""
    ext = os.path.splitext(file_path)[1].lower()
    
    logger.info(f"Processing file: {file_path}, Extension: {ext}")
    
    if ext in ['.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.gif']:
        return extract_text_from_image(file_path)
    elif ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext == '.docx':
        return extract_text_from_docx(file_path)
    else:
        return None, f"Unsupported file type: {ext}"

# =======================
# FLASK ROUTES
# =======================

@app.route('/')
def index():
    """Main page"""
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    """Handle file upload and text extraction"""
    filepath = None
    try:
        # Check if file is present
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': 'No file provided'}), 400
        
        file = request.files['file']
        
        if file.filename == '':
            return jsonify({'success': False, 'error': 'No file selected'}), 400
        
        if not allowed_file(file.filename):
            return jsonify({
                'success': False, 
                'error': f'File type not allowed. Supported: {", ".join(app.config["ALLOWED_EXTENSIONS"])}'
            }), 400
        
        # Save uploaded file
        filename = generate_unique_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        
        logger.info(f"File uploaded: {filename}")
        
        # Extract text
        extracted_text, error = extract_text_from_file(filepath)
        
        if error:
            # Clean up uploaded file
            if filepath and os.path.exists(filepath):
                os.remove(filepath)
            return jsonify({'success': False, 'error': error}), 400
        
        # Save extracted text
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = f"extracted_{timestamp}.txt"
        output_path = os.path.join(app.config['OUTPUT_FOLDER'], output_filename)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(extracted_text)
        
        # Store in session for download
        session['last_extraction'] = output_filename
        
        # Clean up uploaded file
        if filepath and os.path.exists(filepath):
            os.remove(filepath)
        
        logger.info(f"Text extracted successfully: {len(extracted_text)} characters")
        
        return jsonify({
            'success': True,
            'text': extracted_text,
            'filename': output_filename,
            'char_count': len(extracted_text),
            'word_count': len(extracted_text.split())
        })
        
    except Exception as e:
        logger.error(f"Upload error: {e}", exc_info=True)
        # Clean up on error
        if filepath and os.path.exists(filepath):
            try:
                os.remove(filepath)
            except:
                pass
        return jsonify({'success': False, 'error': f'Server error: {str(e)}'}), 500

@app.route('/download/<filename>')
def download_file(filename):
    """Download extracted text file"""
    try:
        filepath = os.path.join(app.config['OUTPUT_FOLDER'], secure_filename(filename))
        
        if not os.path.exists(filepath):
            return jsonify({'success': False, 'error': 'File not found'}), 404
        
        return send_file(filepath, as_attachment=True, download_name=filename)
        
    except Exception as e:
        logger.error(f"Download error: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/health')
def health_check():
    """Health check endpoint"""
    return jsonify({
        'status': 'healthy',
        'timestamp': datetime.now().isoformat(),
        'tesseract_available': check_tesseract()
    })

def check_tesseract():
    """Check if Tesseract is available"""
    try:
        pytesseract.get_tesseract_version()
        return True
    except:
        return False

# =======================
# ERROR HANDLERS
# =======================

@app.errorhandler(413)
def too_large(e):
    """Handle file too large error"""
    return jsonify({
        'success': False,
        'error': 'File too large. Maximum size is 50MB'
    }), 413

@app.errorhandler(404)
def not_found(e):
    """Handle 404 errors"""
    return jsonify({
        'success': False,
        'error': 'Resource not found'
    }), 404

@app.errorhandler(500)
def internal_error(e):
    """Handle internal server errors"""
    logger.error(f"Internal error: {e}")
    return jsonify({
        'success': False,
        'error': 'Internal server error'
    }), 500

# =======================
# RUN APP
# =======================

if __name__ == '__main__':
    # Check dependencies on startup
    if not check_tesseract():
        logger.warning("⚠️  Tesseract OCR not found! Please install it.")
    else:
        logger.info("✓ Tesseract OCR is available")
    
    app.run(debug=True, host='0.0.0.0', port=5000)